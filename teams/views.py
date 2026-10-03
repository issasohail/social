import csv
import re
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from PIL import Image, ImageDraw, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from boards.models import Board
from organization.access import active_accesses
from organization.models import Jamatkhana, LocalCouncil, NationalCouncil, RegionalCouncil
from .forms import TeamAppointmentForm
from .models import Portfolio, TeamAppointment, TeamCategory, TeamPosition, Term


def _digits(value):
    return re.sub(r'\D', '', value or '')


def _scoped(qs, user):
    if user.is_superuser:
        return qs
    allowed = Q()
    for access in active_accesses(user):
        board_q = Q() if not access.board_id else Q(board_id=access.board_id)
        if access.level == 'NATIONAL':
            if access.national_council_id:
                allowed |= board_q & Q(national_council_id=access.national_council_id)
            else:
                allowed |= board_q
        elif access.level == 'REGIONAL' and access.regional_council_id:
            allowed |= board_q & Q(regional_council_id=access.regional_council_id)
        elif access.level == 'LOCAL' and access.local_council_id:
            allowed |= board_q & Q(local_council_id=access.local_council_id)
        elif access.level == 'JK' and access.jamatkhana_id:
            allowed |= board_q & (Q(jamatkhana_id=access.jamatkhana_id) | Q(covered_jamatkhanas=access.jamatkhana_id))
    return qs.filter(allowed).distinct() if allowed else qs.none()


def _filtered(request):
    qs = TeamAppointment.objects.select_related(
        'term','board','person','position','position__category','portfolio','national_council',
        'regional_council','local_council','jamatkhana'
    ).prefetch_related('covered_jamatkhanas__local_council__regional_council')
    qs = _scoped(qs, request.user)
    term = request.GET.get('term') or ''
    if not term:
        current = Term.objects.filter(is_current=True, is_active=True).first()
        if current:
            term = str(current.pk)
    filters = {
        'term': term,
        'board': request.GET.get('board',''),
        'level': request.GET.get('level',''),
        'regional': request.GET.get('regional',''),
        'local': request.GET.get('local',''),
        'jk': request.GET.get('jk',''),
        'portfolio': request.GET.get('portfolio',''),
        'position': request.GET.get('position',''),
        'status': request.GET.get('status','active'),
        'q': request.GET.get('q','').strip(),
    }
    if filters['term']: qs = qs.filter(term_id=filters['term'])
    if filters['board']: qs = qs.filter(board_id=filters['board'])
    if filters['level']: qs = qs.filter(level=filters['level'])
    if filters['regional']: qs = qs.filter(regional_council_id=filters['regional'])
    if filters['local']: qs = qs.filter(local_council_id=filters['local'])
    if filters['jk']: qs = qs.filter(Q(jamatkhana_id=filters['jk'])|Q(covered_jamatkhanas=filters['jk'])).distinct()
    if filters['portfolio']: qs = qs.filter(portfolio_id=filters['portfolio'])
    if filters['position']: qs = qs.filter(position_id=filters['position'])
    if filters['status'] == 'active': qs = qs.filter(is_active=True)
    elif filters['status'] == 'inactive': qs = qs.filter(is_active=False)
    if filters['q']:
        q = filters['q']
        qs = qs.filter(
            Q(person__full_name__icontains=q)|Q(person__mobile__icontains=q)|Q(person__whatsapp_number__icontains=q)|
            Q(person__email__icontains=q)|Q(person__jamatkhana__name__icontains=q)|Q(person__local_council__name__icontains=q)|
            Q(position__name__icontains=q)|Q(portfolio__name__icontains=q)|Q(board__name__icontains=q)|Q(board__short_name__icontains=q)
        )
    return qs.order_by('position__sort_order','position__name','person__full_name'), filters


def _hierarchy_maps():
    regions = RegionalCouncil.objects.filter(is_active=True).select_related('national_council')
    locals_ = LocalCouncil.objects.filter(is_active=True).select_related('regional_council__national_council')
    jks = Jamatkhana.objects.filter(is_active=True).select_related('local_council__regional_council__national_council')
    return {
        'region_map': {str(x.pk): {'national': x.national_council_id} for x in regions},
        'local_map': {str(x.pk): {'regional': x.regional_council_id, 'national': x.regional_council.national_council_id} for x in locals_},
        'jk_map': {str(x.pk): {'local': x.local_council_id, 'regional': x.local_council.regional_council_id, 'national': x.local_council.regional_council.national_council_id} for x in jks},
        'portfolio_map': {str(x.pk): {'board': x.board_id} for x in Portfolio.objects.filter(is_active=True)},
    }


@login_required
def team_list(request):
    qs, filters = _filtered(request)
    selected_board = Board.objects.filter(pk=filters['board']).first() if filters['board'] else None
    ctx = {
        'appointments': qs,
        'filters': filters,
        'selected_board': selected_board,
        'terms': Term.objects.filter(is_active=True),
        'boards': Board.objects.filter(is_active=True).order_by('sort_order','name'),
        'positions': TeamPosition.objects.filter(is_active=True).select_related('category'),
        'portfolios': Portfolio.objects.filter(is_active=True).select_related('board'),
        'regions': RegionalCouncil.objects.filter(is_active=True).select_related('national_council'),
        'locals': LocalCouncil.objects.filter(is_active=True).select_related('regional_council'),
        'jks': Jamatkhana.objects.filter(is_active=True).select_related('local_council__regional_council'),
        'levels': TeamAppointment.Level.choices,
    }
    return render(request, 'teams/list.html', ctx)


@login_required
def team_detail(request, pk):
    obj = get_object_or_404(_scoped(TeamAppointment.objects.select_related('term','board','person','position','portfolio','national_council','regional_council','local_council','jamatkhana').prefetch_related('covered_jamatkhanas'), request.user), pk=pk)
    history = TeamAppointment.objects.filter(person=obj.person).select_related('term','board','position','portfolio').order_by('-term__start_date')
    return render(request, 'teams/detail.html', {'appointment':obj,'history':history})


@login_required
def team_create(request):
    if not (request.user.is_superuser or request.user.has_perm('teams.add_teamappointment')):
        messages.error(request, 'You do not have permission to add team appointments.')
        return redirect('team_list')
    form = TeamAppointmentForm(request.POST or None)
    if form.is_valid():
        obj = form.save(commit=False)
        obj.created_by=request.user
        obj.save()
        form.save_m2m()
        messages.success(request, 'Team appointment added.')
        return redirect('team_detail', pk=obj.pk)
    ctx={'form':form,'title':'Add team appointment', **_hierarchy_maps()}
    return render(request,'teams/form.html',ctx)


@login_required
def team_edit(request, pk):
    obj = get_object_or_404(_scoped(TeamAppointment.objects.all(), request.user), pk=pk)
    if not (request.user.is_superuser or request.user.has_perm('teams.change_teamappointment')):
        messages.error(request, 'You do not have permission to edit team appointments.')
        return redirect('team_detail', pk=pk)
    form=TeamAppointmentForm(request.POST or None, instance=obj)
    if form.is_valid():
        form.save()
        messages.success(request,'Team appointment updated.')
        return redirect('team_detail',pk=pk)
    ctx={'form':form,'title':'Edit team appointment','appointment':obj, **_hierarchy_maps()}
    return render(request,'teams/form.html',ctx)


def _coverage_label(jk):
    # Jurisdiction already carries council hierarchy; coverage only needs the Jamatkhana name.
    return jk.name


def _export_rows(qs):
    for i,a in enumerate(qs,1):
        coverage=', '.join(_coverage_label(x) for x in a.covered_jamatkhanas.all())
        phone=a.person.whatsapp_number or a.person.mobile or ''
        yield [i,a.person.full_name,a.position.name,a.portfolio.name if a.portfolio else '',a.get_level_display(),str(a.jurisdiction_name or ''),coverage,phone,a.person.email,a.term.name,'Active' if a.is_active else 'Inactive']


def _export_title(filters):
    term_id = filters.get('term') or None
    board_id = filters.get('board') or None
    term = Term.objects.filter(pk=term_id).first() if term_id else None
    board = Board.objects.filter(pk=board_id).first() if board_id else None
    board_name=(board.short_name or board.name) if board else 'Teams & Contacts'
    title=f'{board_name} Contact List'
    if term:
        title += f' — {term.name}'
    return title, board


@login_required
def team_export(request, export_format):
    qs, filters = _filtered(request)
    rows=list(_export_rows(qs))
    stamp=timezone.localtime().strftime('%d %b %Y %H:%M')
    title, board = _export_title(filters)
    headers=['S/N','Name','Position','Portfolio','Level','Jurisdiction','JK Coverage','Phone','Email','Term','Status']
    filename_base=((board.code if board else 'teams') + '_contact_list').lower()

    if export_format == 'csv':
        response=HttpResponse(content_type='text/csv; charset=utf-8')
        response['Content-Disposition']=f'attachment; filename="{filename_base}.csv"'
        writer=csv.writer(response)
        writer.writerow([title]); writer.writerow(['Generated',stamp]); writer.writerow([]); writer.writerow(headers)
        writer.writerows(rows)
        return response

    if export_format == 'xlsx':
        wb=Workbook(); ws=wb.active; ws.title='Contact List'
        last_col=len(headers)
        ws.merge_cells(start_row=1,start_column=1,end_row=1,end_column=last_col)
        heading=ws.cell(1,1,title); heading.font=Font(size=16,bold=True); heading.alignment=Alignment(horizontal='center',vertical='center')
        ws.row_dimensions[1].height=28
        ws.merge_cells(start_row=2,start_column=1,end_row=2,end_column=last_col)
        generated=ws.cell(2,1,f'Generated: {stamp}'); generated.font=Font(size=9,italic=True); generated.alignment=Alignment(horizontal='center')
        for c,h in enumerate(headers,1):
            cell=ws.cell(4,c,h); cell.font=Font(bold=True); cell.fill=PatternFill('solid',fgColor='DDEFEF'); cell.alignment=Alignment(horizontal='center',vertical='center')
        for r_index,row in enumerate(rows,5):
            for c_index,value in enumerate(row,1):
                cell=ws.cell(r_index,c_index,value)
                cell.alignment=Alignment(vertical='top',wrap_text=True)
            phone=_digits(row[7])
            if phone:
                cell=ws.cell(r_index,8,row[7])
                cell.hyperlink=f'https://wa.me/{phone}'
                cell.style='Hyperlink'
        widths=[7,24,22,22,16,30,44,18,32,15,12]
        for idx,w in enumerate(widths,1): ws.column_dimensions[ws.cell(1,idx).column_letter].width=w
        ws.freeze_panes='A5'
        ws.sheet_view.showGridLines=False
        ws.oddFooter.left.text=f'Generated: {stamp}'
        ws.oddFooter.right.text='Page &[Page] of &[Pages]'
        ws.evenFooter.left.text=f'Generated: {stamp}'
        ws.evenFooter.right.text='Page &[Page] of &[Pages]'
        ws.print_title_rows='1:4'
        ws.page_setup.orientation='landscape'; ws.page_setup.fitToWidth=1; ws.page_setup.fitToHeight=0
        out=BytesIO(); wb.save(out)
        r=HttpResponse(out.getvalue(),content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        r['Content-Disposition']=f'attachment; filename="{filename_base}.xlsx"'
        return r

    if export_format == 'pdf':
        out=BytesIO(); styles=getSampleStyleSheet()
        pdf_title=f'{(board.code if board else "Teams")} Contact List'
        def footer(canvas,doc):
            canvas.saveState(); canvas.setTitle(pdf_title); canvas.setAuthor('Social Welfare Center')
            canvas.setFont('Helvetica',8); canvas.drawString(12*mm,8*mm,f'Generated: {stamp}'); canvas.drawRightString(285*mm,8*mm,f'Page {doc.page}'); canvas.restoreState()
        doc=SimpleDocTemplate(out,pagesize=landscape(A4),leftMargin=9*mm,rightMargin=9*mm,topMargin=10*mm,bottomMargin=14*mm,title=pdf_title)
        data=[[Paragraph(f'<b>{h}</b>',styles['BodyText']) for h in headers]]
        for row in rows:
            rendered=[]
            for idx,v in enumerate(row):
                text=str(v or '')
                if idx==7 and _digits(text):
                    text=f'<link href="https://wa.me/{_digits(text)}" color="#128C7E"><u>{text}</u></link>'
                rendered.append(Paragraph(text,styles['BodyText']))
            data.append(rendered)
        table=Table(data,repeatRows=1,colWidths=[9*mm,33*mm,28*mm,29*mm,19*mm,31*mm,50*mm,26*mm,37*mm,22*mm,18*mm])
        table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#DDEFEF')),('GRID',(0,0),(-1,-1),0.35,colors.HexColor('#cbd5e1')),('VALIGN',(0,0),(-1,-1),'TOP'),('FONTSIZE',(0,0),(-1,-1),7),('LEFTPADDING',(0,0),(-1,-1),3),('RIGHTPADDING',(0,0),(-1,-1),3)]))
        doc.build([Paragraph(title,styles['Title']),Spacer(1,2*mm),Paragraph(f'Generated {stamp}',styles['Normal']),Spacer(1,4*mm),table],onFirstPage=footer,onLaterPages=footer)
        r=HttpResponse(out.getvalue(),content_type='application/pdf'); r['Content-Disposition']=f'attachment; filename="{filename_base}.pdf"'; return r

    if export_format == 'jpg':
        widths=[50,260,190,190,130,220,340,170,260,120,100]; width=sum(widths)+40; rowh=48; height=max(300,175+(len(rows)+1)*rowh+60)
        img=Image.new('RGB',(width,height),'white'); d=ImageDraw.Draw(img); font=ImageFont.load_default(); d.text((20,20),title,fill='black',font=font); d.text((20,45),f'Generated: {stamp}',fill='black',font=font)
        y=88; x=20
        for j,h in enumerate(headers): d.rectangle((x,y,x+widths[j],y+rowh),outline='gray'); d.text((x+4,y+16),h,fill='black',font=font); x+=widths[j]
        y+=rowh
        for row in rows:
            x=20
            for j,v in enumerate(row):
                d.rectangle((x,y,x+widths[j],y+rowh),outline='lightgray'); txt=str(v or ''); maxchars=max(5,widths[j]//7); d.text((x+4,y+16),txt[:maxchars],fill='black',font=font); x+=widths[j]
            y+=rowh
        d.text((20,height-25),f'Generated: {stamp}',fill='black',font=font); d.text((width-90,height-25),'Page 1',fill='black',font=font)
        out=BytesIO(); img.save(out,'JPEG',quality=90)
        r=HttpResponse(out.getvalue(),content_type='image/jpeg'); r['Content-Disposition']=f'attachment; filename="{filename_base}.jpg"'; return r
    return HttpResponse('Unsupported format',status=400)


@login_required
def term_list(request):
    if not request.user.is_superuser:
        return redirect('team_list')
    return render(request,'teams/terms.html',{'terms':Term.objects.all()})


@login_required
def team_settings(request):
    if not request.user.is_superuser:
        return redirect('team_list')
    from .forms import PortfolioForm, TeamCategoryForm, TeamPositionForm, TermForm
    kind=request.GET.get('kind','term')
    forms={'term':TermForm,'category':TeamCategoryForm,'position':TeamPositionForm,'portfolio':PortfolioForm}
    models={'term':Term,'category':TeamCategory,'position':TeamPosition,'portfolio':Portfolio}
    Form=forms.get(kind,TermForm); Model=models.get(kind,Term)
    instance=None
    edit=request.GET.get('edit')
    if edit: instance=get_object_or_404(Model,pk=edit)
    form=Form(request.POST or None,instance=instance)
    if request.method=='POST' and form.is_valid():
        form.save(); messages.success(request,'Team setting saved.'); return redirect(f"{reverse('team_settings')}?kind={kind}")
    items=Model.objects.all()
    if kind=='portfolio': items=items.select_related('board')
    if kind=='position': items=items.select_related('category')
    return render(request,'teams/settings.html',{'kind':kind,'form':form,'items':items})
